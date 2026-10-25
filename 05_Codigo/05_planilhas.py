"""
Planilhas editáveis (Excel) com FÓRMULAS - materialidade, amostragem, analíticos, conciliações,
estimativas, continuidade, distorções, componentes, matriz de cobertura e fontes.
Uso: python 05_Codigo/05_planilhas.py      Saída: 03_Planilhas/*.xlsx
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import notas_dfp2025 as ne  # noqa: E402
from cobertura import FONTES, MATRIZ  # noqa: E402
from parametros import (LIMITE_TRIVIAL, MATERIALIDADE_ESPECIFICA, PREPARADOR, DATA_PREPARO, RESULTADOS,  # noqa: E402
                        bases_materialidade, conta)

OUT = Path(__file__).resolve().parents[1] / "03_Planilhas"
CAB = PatternFill("solid", fgColor="1F3A5F")
ENT = PatternFill("solid", fgColor="FFF4CC")
SINT = PatternFill("solid", fgColor="FDE2E1")
B = Font(bold=True)
BR = Font(bold=True, color="FFFFFF")
FMT = '#,##0;[Red]-#,##0'
FMT2 = '#,##0.00;[Red]-#,##0.00'
PCT = '0.00%'


def cabecalho(ws, titulo, norma, linha_ini=1, sintetico=False):
    ws.cell(linha_ini, 1, titulo).font = Font(bold=True, size=13)
    ws.cell(linha_ini + 1, 1, f"Norma: {norma} | Preparador: {PREPARADOR} | Data: {DATA_PREPARO} | Revisão: ________ (campo em branco)")
    ws.cell(linha_ini + 2, 1, "Células amarelas = premissas editáveis; demais valores são fórmulas." +
            ("  *** DADOS SINTÉTICOS PARA FINS DIDÁTICOS ***" if sintetico else "  Fonte: CVM Dados Abertos / DFP 2025."))
    if sintetico:
        ws.cell(linha_ini + 2, 1).fill = SINT
    return linha_ini + 4


def tabela(ws, linha, colunas, larguras=None):
    for j, c in enumerate(colunas, 1):
        cel = ws.cell(linha, j, c)
        cel.font, cel.fill = BR, CAB
        cel.alignment = Alignment(wrap_text=True, vertical="center")
        if larguras:
            ws.column_dimensions[get_column_letter(j)].width = larguras[j - 1]
    return linha + 1


def df_para_ws(ws, df, linha, fmt=None, larguras=None):
    linha = tabela(ws, linha, list(df.columns), larguras)
    for r in df.itertuples(index=False):
        for j, v in enumerate(r, 1):
            cel = ws.cell(linha, j, None if (isinstance(v, float) and pd.isna(v)) else v)
            if isinstance(v, (int, float)) and fmt:
                cel.number_format = fmt
        linha += 1
    return linha


# ----------------------------------------------------------------------------------------- PT 320
def pt320():
    wb = Workbook()
    ws = wb.active
    ws.title = "PT320_Materialidade"
    l = cabecalho(ws, "PT 320 - Materialidade no planejamento e na execução", "NBC TA 320 (R1), itens 10-14, A3-A13")
    l = tabela(ws, l, ["Base de referência", "Valor 2025 (R$ mil)", "% mínimo", "% máximo", "Faixa mínima", "Faixa máxima",
                       "% aplicado", "Materialidade calculada", "Escolhida?"], [62, 18, 10, 10, 16, 16, 11, 20, 11])
    ini = l
    for b in bases_materialidade():
        ws.cell(l, 1, b["base"])
        for j, k in [(2, "valor"), (3, "pct_min"), (4, "pct_max"), (7, "pct")]:
            c = ws.cell(l, j, b[k])
            c.fill = ENT
            c.number_format = FMT if j == 2 else PCT
        ws.cell(l, 5, f"=B{l}*C{l}").number_format = FMT
        ws.cell(l, 6, f"=B{l}*D{l}").number_format = FMT
        ws.cell(l, 8, f"=B{l}*G{l}").number_format = FMT
        ws.cell(l, 9, "SIM" if b["escolhida"] else "não")
        l += 1
    l += 1
    linhas = [("Materialidade global calculada (base escolhida)", f"=INDEX(H{ini}:H{l-2},MATCH(\"SIM\",I{ini}:I{l-2},0))", FMT),
              ("Materialidade global adotada (arredondada para baixo)", f"=ROUNDDOWN(B{l}/5000,0)*5000", FMT),
              ("% para materialidade de execução (auditoria inicial, aquisições, 1º ano do auditor)", 0.60, PCT),
              ("Materialidade para execução", f"=B{l+1}*B{l+2}", FMT),
              ("% limite de distorções claramente triviais", 0.05, PCT),
              ("Limite claramente trivial", f"=B{l+1}*B{l+4}", FMT),
              ("Distorção tolerável para testes do imobilizado (= execução)", f"=B{l+3}", FMT)]
    for i, (rot, val, f) in enumerate(linhas):
        ws.cell(l + i, 1, rot).font = B
        c = ws.cell(l + i, 2, val)
        c.number_format = f
        if not str(val).startswith("="):
            c.fill = ENT
    l += len(linhas) + 1
    ws.cell(l, 1, "Materialidades específicas (NBC TA 320 item 10)").font = B
    l += 1
    for k, v in MATERIALIDADE_ESPECIFICA.items():
        ws.cell(l, 1, k)
        c = ws.cell(l, 2, v)
        c.number_format, c.fill = FMT, ENT
        l += 1
    l += 1
    for t in ["Justificativa: a receita é estável e é o principal indicador para usuários de uma produtora de petróleo; o LAIR de 2025 é",
              "negativo (R$ -232.613 mil) e distorcido por variação cambial e IR diferido, portanto inadequado como base (NBC TA 320 A4-A8).",
              "Validação cruzada: 2% do EBITDA ajustado e 0,25% do ativo total ficam na faixa de R$ 160 mi, o que corrobora o valor adotado.",
              "Revisão: reavaliar se o fechamento dos 20% de Peregrino ou mudança relevante no Brent alterar as bases (item 12)."]:
        ws.cell(l, 1, t)
        l += 1
    wb.save(OUT / "PT320_Materialidade.xlsx")


# ----------------------------------------------------------------------------------------- PT 530
def pt530():
    r = json.loads((RESULTADOS / "resumo_testes.json").read_text(encoding="utf-8"))
    m = r["mus"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Parametros_MUS"
    l = cabecalho(ws, "PT 530 - Amostragem por unidade monetária (MUS) - Imobilizado", "NBC TA 530 itens 6-15, A10-A23", sintetico=True)
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 22
    par = [("Valor contábil da população (itens > 0) R$", m["valor_contabil_populacao_brl"], FMT2, True),
           ("Itens na população", m["populacao_itens"], FMT, True),
           ("Nível de confiança", 0.95, PCT, True),
           ("Fator de confiabilidade (0 erros, Poisson 95%)", 3.0, "0.00", True),
           ("Distorção tolerável R$ (materialidade de execução)", 93_000_000, FMT2, True),
           ("Distorção esperada (% da tolerável)", 0.10, PCT, True),
           ("Fator de expansão", 1.6, "0.00", True)]
    for i, (rot, v, f, ent_) in enumerate(par):
        ws.cell(l + i, 1, rot)
        c = ws.cell(l + i, 2, v)
        c.number_format = f
        if ent_:
            c.fill = ENT
    a = l
    l += len(par)
    calc = [("Distorção esperada R$", f"=B{a+4}*B{a+5}", FMT2),
            ("Tamanho da amostra n = VC x FC / (DT - DE x FE)", f"=ROUNDUP(B{a}*B{a+3}/(B{a+4}-B{l}*B{a+6}),0)", FMT),
            ("Intervalo amostral J = VC / n", f"=B{a}/B{l+1}", FMT2),
            ("Início aleatório (semente 231019943) R$", m["inicio_aleatorio_brl"], FMT2),
            ("Itens distintos selecionados (seleção sistemática PPS)", m["itens_selecionados"], FMT),
            ("Itens do estrato topo (>= J) - 100% testados", m["itens_topo"], FMT)]
    for i, (rot, v, f) in enumerate(calc):
        ws.cell(l + i, 1, rot).font = B
        ws.cell(l + i, 2, v).number_format = f
    l += len(calc) + 1
    ws.cell(l, 1, "Avaliação - ver aba Avaliacao_MUS (distorções ordenadas por taint, fatores incrementais)").font = B

    ws2 = wb.create_sheet("Amostra_MUS")
    am = pd.read_csv(RESULTADOS / "PT530_amostra_mus.csv")
    cols = ["id_registro", "classe_ne", "entidade", "campo_ugc", "descricao", "estrato", "valor_contabil_brl", "valor_auditado_brl"]
    l2 = cabecalho(ws2, "Amostra selecionada e resultado do teste (valor auditado = documentação-suporte simulada)", "NBC TA 530 item 9-11", sintetico=True)
    l2 = tabela(ws2, l2, cols + ["distorcao_brl", "taint"], [12, 34, 12, 14, 34, 18, 18, 18, 16, 8])
    for rr in am[cols].itertuples(index=False):
        for j, v in enumerate(rr, 1):
            ws2.cell(l2, j, v)
        ws2.cell(l2, 7).number_format = ws2.cell(l2, 8).number_format = FMT2
        ws2.cell(l2, 9, f"=G{l2}-H{l2}").number_format = FMT2
        ws2.cell(l2, 10, f"=IF(G{l2}=0,0,MIN(1,I{l2}/G{l2}))").number_format = "0.0000"
        l2 += 1

    ws3 = wb.create_sheet("Avaliacao_MUS")
    l3 = cabecalho(ws3, "Avaliação: projeção e limite superior de distorção (Stringer / Poisson 95%)", "NBC TA 530 itens 14-15, A21-A23", sintetico=True)
    ws3.cell(l3, 1, "Intervalo J")
    ws3.cell(l3, 2, "=Parametros_MUS!B14").number_format = FMT2
    jrow = l3
    l3 = tabela(ws3, l3 + 2, ["id_registro", "estrato", "valor_contabil", "distorcao", "taint", "fator k", "fator k-1",
                              "incremento (fk - fk-1 - 1)", "distorção projetada", "provisão incremental"], [12, 18, 16, 16, 8, 8, 8, 12, 18, 18])
    fatores = [3.00, 4.75, 6.30, 7.76, 9.16, 10.52, 11.85, 13.15, 14.44, 15.71, 16.97, 18.21, 19.45, 20.67, 21.89]
    d = am[am["distorcao_brl"] > 0.01].copy()
    d["topo"] = d["estrato"].str.startswith("Topo")
    d = pd.concat([d[~d["topo"]].sort_values("taint", ascending=False), d[d["topo"]]])
    k = 0
    pri = l3
    for rr in d.itertuples(index=False):
        ws3.cell(l3, 1, rr.id_registro)
        ws3.cell(l3, 2, rr.estrato)
        ws3.cell(l3, 3, rr.valor_contabil_brl).number_format = FMT2
        ws3.cell(l3, 4, rr.distorcao_brl).number_format = FMT2
        ws3.cell(l3, 5, f"=MIN(1,D{l3}/C{l3})").number_format = "0.0000"
        if not rr.topo:
            k += 1
            ws3.cell(l3, 6, fatores[k] if k < len(fatores) else fatores[-1])
            ws3.cell(l3, 7, fatores[k - 1] if k - 1 < len(fatores) else fatores[-1])
            ws3.cell(l3, 8, f"=F{l3}-G{l3}-1")
            ws3.cell(l3, 9, f"=E{l3}*$B${jrow}").number_format = FMT2
            ws3.cell(l3, 10, f"=H{l3}*E{l3}*$B${jrow}").number_format = FMT2
        else:
            ws3.cell(l3, 9, f"=D{l3}").number_format = FMT2
            ws3.cell(l3, 10, 0)
        l3 += 1
    ult = l3 - 1
    res = [("Precisão básica = FC x J", f"=3*$B${jrow}"), ("Distorção projetada", f"=SUM(I{pri}:I{ult})"),
           ("Provisão incremental", f"=SUM(J{pri}:J{ult})"), ("Limite superior de distorção (LSD)", f"=B{l3+1}+B{l3+2}+B{l3+3}"),
           ("Distorção tolerável", "=Parametros_MUS!B9"), ("Conclusão", f'=IF(B{l3+4}>B{l3+5},"Rejeitar: LSD > tolerável","Aceitar")')]
    for i, (rot, f) in enumerate(res, 1):
        ws3.cell(l3 + i, 1, rot).font = B
        ws3.cell(l3 + i, 2, f).number_format = FMT2

    ws4 = wb.create_sheet("Atributos_Controle")
    at = r["atributos"]
    l4 = cabecalho(ws4, "Amostragem por atributos - controle de aprovação (segregação e alçada)", "NBC TA 530 / NBC TA 330 itens 8-17", sintetico=True)
    ws4.column_dimensions["A"].width = 55
    for i, (rot, v, f) in enumerate([("População", at["populacao"], FMT), ("Confiança", 0.95, PCT), ("Taxa tolerável de desvio", 0.05, PCT),
                                     ("Taxa esperada de desvio", 0.005, PCT), ("Tamanho da amostra (tabela AICPA)", 93, FMT),
                                     ("Desvios encontrados", at["desvios_encontrados"], FMT)]):
        ws4.cell(l4 + i, 1, rot)
        c = ws4.cell(l4 + i, 2, v)
        c.number_format, c.fill = f, ENT
    ws4.cell(l4 + 6, 1, "Limite superior de desvio = fator Poisson(k)/n").font = B
    ws4.cell(l4 + 6, 2, f"=IF(B{l4+5}=0,3,IF(B{l4+5}=1,4.75,IF(B{l4+5}=2,6.3,IF(B{l4+5}=3,7.76,9.16))))/B{l4+4}").number_format = PCT
    ws4.cell(l4 + 7, 1, "Conclusão").font = B
    ws4.cell(l4 + 7, 2, f'=IF(B{l4+6}<=B{l4+2},"Controle efetivo","Controle não efetivo")')
    wb.save(OUT / "PT530_Amostragem.xlsx")


# ----------------------------------------------------------------------------------------- PT 520
def pt520():
    rm = json.loads((RESULTADOS / "resumo_analiticos.json").read_text(encoding="utf-8"))["mercado"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Expectativas"
    l = cabecalho(ws, "PT 520 - Procedimentos analíticos substantivos (cálculos reproduzíveis)", "NBC TA 520 itens 5-7, A12-A16")
    ws.column_dimensions["A"].width = 64
    ws.column_dimensions["B"].width = 20
    entradas = [("Brent médio 2025 US$/bbl (EIA)", rm["brent_medio_2025"]), ("Brent médio 2024 US$/bbl (EIA)", rm["brent_medio_2024"]),
                ("PTAX média 2025 R$/US$ (BCB SGS 1)", rm["ptax_media_2025"]), ("PTAX média 2024 (BCB)", rm["ptax_media_2024"]),
                ("Volume vendido 2025 - milhões bbl (RA p. 27)", 37.8), ("Desconto de qualidade premissa US$/bbl", 4.0),
                ("Receita bruta de óleo próprio registrada R$ mil (NE 22)", sum(x[1] for x in ne.NE22_RECEITA if x[0] != "Trading")),
                ("Dívida bruta 31/12/2024 R$ mil (NE 13+14)", 11_936_077 + 6_465_871), ("Dívida bruta 31/12/2025 R$ mil (NE 13+14)", 17_167_620 + 9_943_197),
                ("Custo médio da dívida (RA p. 39)", 0.0636), ("Juros sobre empréstimos registrados R$ mil (NE 26)", 1_344_905),
                ("Provisão abandono 31/12/2024 R$ mil (NE 17)", 3_387_766), ("Complemento ARO Peregrino 11/11/2025 R$ mil (NE 9.d)", 1_322_571),
                ("Taxa média de desconto ARO (NE 17)", 0.0755), ("Atualização monetária ARO registrada R$ mil (NE 26)", 371_394),
                ("Royalties e PE registrados R$ mil (NE 23)", 1_533_652), ("Alíquota de royalties ANP", 0.10)]
    e = {}
    for i, (rot, v) in enumerate(entradas):
        ws.cell(l + i, 1, rot)
        c = ws.cell(l + i, 2, v)
        c.fill = ENT
        c.number_format = PCT if v < 1 else FMT2
        e[i] = f"B{l+i}"
    l += len(entradas) + 1
    l = tabela(ws, l, ["Conta", "Valor esperado", "Valor registrado", "Diferença", "Diferença %", "Precisão %", "Limite de investigação",
                       "Investigar?", "Precisão isolada suficiente (limite <= ME 93.000)?"], [64, 20, 20, 18, 12, 11, 20, 12, 22])
    linhas = [("Receita bruta de óleo próprio", f"={e[4]}*1000000*({e[0]}-{e[5]})*{e[2]}/1000", f"={e[6]}", 0.05),
              ("Juros sobre empréstimos e debêntures", f"=({e[7]}+{e[8]})/2*{e[9]}", f"={e[10]}", 0.10),
              ("Atualização monetária da provisão de abandono", f"=({e[11]}+{e[12]}*50/365)*{e[13]}", f"={e[14]}", 0.10),
              ("Royalties e participação especial", f"={e[16]}*{e[6]}", f"={e[15]}", 0.15)]
    for (rot, esp, reg, p) in linhas:
        ws.cell(l, 1, rot)
        ws.cell(l, 2, esp).number_format = FMT
        ws.cell(l, 3, reg).number_format = FMT
        ws.cell(l, 4, f"=C{l}-B{l}").number_format = FMT
        ws.cell(l, 5, f"=D{l}/B{l}").number_format = PCT
        c = ws.cell(l, 6, p)
        c.number_format, c.fill = PCT, ENT
        ws.cell(l, 7, f"=ABS(B{l})*F{l}").number_format = FMT
        ws.cell(l, 8, f'=IF(ABS(D{l})>G{l},"SIM","não")')
        ws.cell(l, 9, f'=IF(G{l}<=93000,"sim","não - combinar com teste de detalhe")')
        l += 1
    ws.cell(l + 1, 1, "D&A: expectativa e decomposição por ativo na aba DA_Decomposicao (investigação da diferença).").font = B
    an = pd.read_csv(RESULTADOS / "PT520_procedimentos_analiticos.csv")
    df_para_ws(ws, an, l + 3, FMT2)

    ws2 = wb.create_sheet("DA_Decomposicao")
    l2 = cabecalho(ws2, "Investigação da diferença de D&A - decomposição das movimentações (NE 8, 9 e 15)", "NBC TA 520 item 7")
    df_para_ws(ws2, pd.read_csv(RESULTADOS / "PT520_decomposicao_da.csv"), l2, FMT, [70, 20])

    ws3 = wb.create_sheet("Horizontal_Vertical")
    l3 = cabecalho(ws3, "Análise horizontal e vertical (BP e DRE consolidados)", "NBC TA 520 / NBC TA 315 A27")
    hv = pd.read_csv(RESULTADOS / "PT520_analise_horizontal_vertical.csv", dtype={"CD_CONTA": str})
    l3 = tabela(ws3, l3, ["Demonstração", "Conta", "Descrição", "2024", "2025", "Var. R$ mil", "Var. %", "AV 2025",
                          "Acima do limite (> ME e > 10%)?"], [10, 12, 50, 16, 16, 16, 10, 10, 16])
    tot = {}
    for rr in hv.itertuples(index=False):
        ws3.cell(l3, 1, rr.demonstracao)
        ws3.cell(l3, 2, rr.CD_CONTA)
        ws3.cell(l3, 3, rr.DS_CONTA)
        ws3.cell(l3, 4, rr._3).number_format = FMT
        ws3.cell(l3, 5, rr._4).number_format = FMT
        if rr.CD_CONTA in ("1", "2", "3.01"):
            tot[rr.demonstracao] = l3
        ws3.cell(l3, 6, f"=E{l3}-D{l3}").number_format = FMT
        ws3.cell(l3, 7, f'=IF(D{l3}=0,"",F{l3}/ABS(D{l3}))').number_format = PCT
        ws3.cell(l3, 8, f"=E{l3}/E${tot.get(rr.demonstracao, l3)}").number_format = PCT
        ws3.cell(l3, 9, f'=IF(AND(ABS(F{l3})>93000,ABS(IF(D{l3}=0,1,F{l3}/D{l3}))>0.1),"SIM","")')
        l3 += 1
    ws4 = wb.create_sheet("Benford_SINTETICO")
    l4 = cabecalho(ws4, "Teste de Benford - primeiro dígito do custo (base sintética)", "NBC TA 240 A44 / NBC TA 520", sintetico=True)
    bf = pd.read_csv(RESULTADOS / "PT520_benford_custo.csv")
    l4 = tabela(ws4, l4, ["Dígito", "Observado", "Esperado = LOG10(1+1/d)", "Desvio absoluto"], [10, 14, 22, 16])
    p0 = l4
    for rr in bf.itertuples(index=False):
        ws4.cell(l4, 1, int(rr.digito))
        ws4.cell(l4, 2, rr.observado).number_format = PCT
        ws4.cell(l4, 3, f"=LOG10(1+1/A{l4})").number_format = PCT
        ws4.cell(l4, 4, f"=ABS(B{l4}-C{l4})").number_format = "0.0000"
        l4 += 1
    ws4.cell(l4, 3, "MAD (conformidade próxima < 0,006)").font = B
    ws4.cell(l4, 4, f"=AVERAGE(D{p0}:D{l4-1})").number_format = "0.0000"
    wb.save(OUT / "PT520_Analiticos.xlsx")


# ------------------------------------------------------------------------------ PT 500 / PT 510
def pt500():
    wb = Workbook()
    ws = wb.active
    ws.title = "Notas_x_Demonstracoes"
    l = cabecalho(ws, "PT 500 - Reconciliação das notas explicativas com as demonstrações (R$ mil)", "NBC TA 500 (R1) itens 6-11, A31")
    rec = pd.read_csv(RESULTADOS / "PT500_reconciliacao_notas_x_demonstracoes.csv")
    l = tabela(ws, l, ["Item", "Ref. demonstração", "Valor demonstração", "Ref. nota", "Valor nota", "Diferença", "Situação", "Explicação"],
               [46, 26, 18, 24, 18, 14, 30, 80])
    for rr in rec.itertuples(index=False):
        ws.cell(l, 1, rr.item)
        ws.cell(l, 2, rr.ref_demonstracao)
        ws.cell(l, 3, rr.valor_demonstracao).number_format = FMT
        ws.cell(l, 4, rr.ref_nota)
        ws.cell(l, 5, rr.valor_nota).number_format = FMT
        ws.cell(l, 6, f"=C{l}-E{l}").number_format = FMT
        ws.cell(l, 7, f'=IF(ABS(F{l})<1,"Conciliado",IF(ABS(F{l})<={LIMITE_TRIVIAL},"Trivial","Investigar"))')
        ws.cell(l, 8, rr.explicacao if isinstance(rr.explicacao, str) else "")
        l += 1
    l += 1
    ws.cell(l, 1, "Provisões cíveis: BP 2.02.04.01.04 x abertura da NE 30 (p. 96-97)").font = B
    itens = [("BP - Provisões cíveis 2025", 642_170), ("NE 30 - cíveis/regulatórias prováveis", 352_598),
             ("NE 30 - passivo contingente PPA Forte", 221_458), ("NE 30 - passivo contingente PPA Stellina", 25_111)]
    for i, (rot, v) in enumerate(itens, 1):
        ws.cell(l + i, 1, rot)
        ws.cell(l + i, 2, v).number_format = FMT
    ws.cell(l + 5, 1, "Diferença não explicada na nota (2024 conciliava: 352.383 + 221.883 + 76.654 = 650.920)").font = B
    ws.cell(l + 5, 2, f"=B{l+1}-B{l+2}-B{l+3}-B{l+4}").number_format = FMT

    ws2 = wb.create_sheet("Base_SINTETICA_x_NE8")
    l2 = cabecalho(ws2, "Conciliação da base simulada com NE 8/NE 15 por classe (R$ mil)", "NBC TA 500 / NBC TA 530", sintetico=True)
    c = pd.read_csv(RESULTADOS / "PT500_conciliacao_base_x_notas.csv", index_col=0)
    c = c.drop(index="TOTAL")
    l2 = tabela(ws2, l2, ["Classe", "Saldo nota", "Saldo base", "Diferença", "Explicada por erros inseridos", "Não explicada"],
                [50, 18, 18, 16, 22, 16])
    p0 = l2
    for nome, rr in c.iterrows():
        ws2.cell(l2, 1, nome)
        ws2.cell(l2, 2, rr["saldo_nota_rmil"]).number_format = FMT2
        ws2.cell(l2, 3, rr["saldo_base_rmil"]).number_format = FMT2
        ws2.cell(l2, 4, f"=C{l2}-B{l2}").number_format = FMT2
        ws2.cell(l2, 5, rr["explicado_erros_inseridos_rmil"]).number_format = FMT2
        ws2.cell(l2, 6, f"=ROUND(D{l2}-E{l2},2)").number_format = FMT2
        l2 += 1
    ws2.cell(l2, 1, "TOTAL (BP 1.02.03 = 31.142.804)").font = B
    for col in "BCDEF":
        ws2[f"{col}{l2}"] = f"=SUM({col}{p0}:{col}{l2-1})"
        ws2[f"{col}{l2}"].number_format = FMT2

    ws3 = wb.create_sheet("PT510_Saldos_Iniciais")
    l3 = cabecalho(ws3, "PT 510 - Saldos de abertura: DFP 2024 original x comparativo da DFP 2025 (R$ mil)", "NBC TA 510 (R1) itens 5-9")
    si = pd.read_csv(RESULTADOS / "PT510_saldos_iniciais.csv", dtype={"cd_conta": str})
    l3 = tabela(ws3, l3, ["Demonstração", "Conta", "Descrição", "DFP 2024 original", "Comparativo DFP 2025", "Diferença", "Situação"],
                [10, 14, 50, 18, 18, 14, 30])
    for rr in si.itertuples(index=False):
        ws3.cell(l3, 1, rr.demonstracao)
        ws3.cell(l3, 2, rr.cd_conta)
        ws3.cell(l3, 3, rr.ds_conta)
        ws3.cell(l3, 4, rr.dfp2024_original).number_format = FMT
        ws3.cell(l3, 5, rr.comparativo_dfp2025).number_format = FMT
        ws3.cell(l3, 6, f"=E{l3}-D{l3}").number_format = FMT
        ws3.cell(l3, 7, f'=IF(ABS(F{l3})<1,"Igual","Examinar reclassificação")')
        l3 += 1
    wb.save(OUT / "PT500_PT510_Conciliacoes.xlsx")


# ------------------------------------------------------------------------------ PT 540 / PT 570
def pt540_570():
    wb = Workbook()
    ws = wb.active
    ws.title = "PT540_ARO"
    l = cabecalho(ws, "PT 540 - Sensibilidade da provisão para abandono (NE 17) - R$ mil", "NBC TA 540 (R2) itens 13, 18, 23-29, A114-A120")
    ws.cell(l, 1, "Choque de taxa (p.p.)")
    ws.cell(l, 2, 0.01).fill = ENT
    ws.cell(l + 1, 1, "Choque de custo (%)")
    ws.cell(l + 1, 2, 0.10).fill = ENT
    ws.cell(l + 2, 1, "Choque de prazo (anos)")
    ws.cell(l + 2, 2, 2).fill = ENT
    ch = l
    l = tabela(ws, l + 4, ["Campo", "Saldo NE 17", "Ano abandono", "Taxa", "Anos (t)", "Custo futuro implícito = VP x (1+r)^t",
                           "VP taxa -choque", "VP taxa +choque", "VP custo +choque", "VP abandono antecipado", "VP abandono postergado"],
               [16, 14, 10, 8, 8, 22, 14, 14, 14, 16, 16])
    p0 = l
    for campo, saldo, _s, ano, tx in ne.NE17_ARO:
        if saldo <= 0:
            continue
        ws.cell(l, 1, campo)
        ws.cell(l, 2, saldo).number_format = FMT
        ws.cell(l, 3, ano).fill = ENT
        ws.cell(l, 4, tx).number_format = PCT
        ws.cell(l, 5, f"=MAX(C{l}-2025,0.5)")
        ws.cell(l, 6, f"=B{l}*(1+D{l})^E{l}").number_format = FMT
        ws.cell(l, 7, f"=F{l}/(1+D{l}-$B${ch})^E{l}").number_format = FMT
        ws.cell(l, 8, f"=F{l}/(1+D{l}+$B${ch})^E{l}").number_format = FMT
        ws.cell(l, 9, f"=F{l}*(1+$B${ch+1})/(1+D{l})^E{l}").number_format = FMT
        ws.cell(l, 10, f"=F{l}/(1+D{l})^MAX(E{l}-$B${ch+2},0.5)").number_format = FMT
        ws.cell(l, 11, f"=F{l}/(1+D{l})^(E{l}+$B${ch+2})").number_format = FMT
        l += 1
    ws.cell(l, 1, "TOTAL").font = B
    for col in "BFGHIJK":
        ws[f"{col}{l}"] = f"=SUM({col}{p0}:{col}{l-1})"
        ws[f"{col}{l}"].number_format = FMT
    ws.cell(l + 1, 1, "Variação x saldo").font = B
    for col in "GHIJK":
        ws[f"{col}{l+1}"] = f"={col}{l}-$B${l}"
        ws[f"{col}{l+1}"].number_format = FMT
    ws.cell(l + 2, 1, "Variação / MG (155.000)").font = B
    for col in "GHIJK":
        ws[f"{col}{l+2}"] = f"={col}{l+1}/155000"
        ws[f"{col}{l+2}"].number_format = "0.0x"
    ws.cell(l + 4, 1, "Polvo excluído: saldo líquido do fundo de abandono (R$ -11.099 mil); prazo de Tubarão Azul = premissa didática.")

    ws2 = wb.create_sheet("PT570_Indicadores")
    l2 = cabecalho(ws2, "PT 570 - Liquidez, endividamento e covenants (R$ mil)", "NBC TA 570 itens 10-16, A3")
    ws2.column_dimensions["A"].width = 62
    ws2.column_dimensions["B"].width = 18
    ent = [("Ativo circulante (BP 1.01)", conta("BPA", "1.01")), ("Passivo circulante (BP 2.01)", conta("BPP", "2.01")),
           ("Estoques (BP 1.01.04)", conta("BPA", "1.01.04")), ("Caixa e equivalentes (BP 1.01.01)", conta("BPA", "1.01.01")),
           ("Empréstimos (NE 13)", 17_167_620), ("Debêntures sem MTM (NE 14)", 9_943_197),
           ("EBIT (DRE 3.05)", conta("DRE", "3.05")), ("D&A (DFC 6.01.01.02)", conta("DFC_MI", "6.01.01.02")),
           ("Outras despesas líquidas não recorrentes (NE 25)", 491_868),
           ("Pró-forma: produção Peregrino +40% (bpd) - premissa", 40_000), ("Pró-forma: margem EBITDA US$/bbl - premissa", 35),
           ("PTAX média 2025", 5.5855), ("Juros pagos (DFC 6.03.04 + 6.03.13)", 569_515 + 301_069),
           ("FCO (DFC 6.01)", conta("DFC_MI", "6.01")), ("Capex (DFC 6.02.01)", -conta("DFC_MI", "6.02.01")), ("Limite covenant", 2.5)]
    for i, (rot, v) in enumerate(ent):
        ws2.cell(l2 + i, 1, rot)
        c = ws2.cell(l2 + i, 2, v)
        c.fill, c.number_format = ENT, FMT2
    a = l2
    l2 += len(ent) + 1
    calc = [("Liquidez corrente", f"=B{a}/B{a+1}", "0.00"), ("Liquidez seca", f"=(B{a}-B{a+2})/B{a+1}", "0.00"),
            ("Dívida líquida", f"=B{a+4}+B{a+5}-B{a+3}", FMT), ("EBITDA ajustado", f"=B{a+6}+B{a+7}+B{a+8}", FMT),
            ("Dívida líquida / EBITDA ajustado (sem pró-forma)", f"=B{l2+2}/B{l2+3}", "0.00x"),
            ("EBITDA pró-forma Peregrino 40% (LTM)", f"=B{a+9}*365*B{a+10}*B{a+11}/1000", FMT),
            ("Dívida líquida / EBITDA pró-forma", f"=B{l2+2}/(B{l2+3}+B{l2+5})", "0.00x"),
            ("EBITDA pró-forma mínimo para atender o covenant", f"=B{l2+2}/B{a+15}-B{l2+3}", FMT),
            ("Cobertura de juros (EBITDA aj. / juros pagos)", f"=B{l2+3}/B{a+12}", "0.0x"),
            ("FCO - capex", f"=B{a+13}-B{a+14}", FMT),
            ("Situação covenant (pró-forma)", f'=IF(B{l2+6}<=B{a+15},"Atende (depende do pró-forma)","Descumpre")', "@")]
    for i, (rot, f, fm) in enumerate(calc):
        ws2.cell(l2 + i, 1, rot).font = B
        ws2.cell(l2 + i, 2, f).number_format = fm
    ws2.cell(l2 + len(calc) + 1, 1, "Evidência posterior: apresentação 2T26 (05/08/2026) divulga alavancagem de 1,5x e quitação do bond 2026 (US$ 168,7 mi).")
    wb.save(OUT / "PT540_PT570_Estimativas_Continuidade.xlsx")


# ------------------------------------------------------------------------------ PT 450 / PT 600
def pt450_600():
    r = json.loads((RESULTADOS / "resumo_testes.json").read_text(encoding="utf-8"))
    wb = Workbook()
    ws = wb.active
    ws.title = "PT450_Mapa"
    l = cabecalho(ws, "PT 450 - Mapa de distorções identificadas (públicas) e simuladas (R$ mil)", "NBC TA 450 (R1) itens 5-13, A1-A23")
    ws.cell(l, 1, "MG")
    ws.cell(l, 2, 155_000)
    ws.cell(l, 3, "ME")
    ws.cell(l, 4, 93_000)
    ws.cell(l, 5, "Trivial")
    ws.cell(l, 6, LIMITE_TRIVIAL)
    l = tabela(ws, l + 2, ["Ref.", "Origem", "Descrição", "Conta/divulgação", "Tipo (factual/julgamento/projetada)", "Ajuste proposto resultado (R$ mil)",
                           "Ajuste proposto ativo/PL (R$ mil)", "Situação", "Avaliação qualitativa"], [7, 22, 60, 26, 20, 14, 14, 26, 60])
    imp = pd.read_csv(RESULTADOS / "PT450_impacto_erros_inseridos.csv").set_index("tipo_erro")["impacto_valor_contabil_brl"] / 1000
    linhas = [
        ("D01", "Pública - DFP", "Divulgação das provisões cíveis na NE 30 não fecha com o BP: diferença de R$ 43.003 mil", "NE 30 / BP 2.02.04.01.04",
         "Factual (divulgação)", 0, 0, "Não corrigida (comunicar)", "Inconsistência de divulgação acima do trivial; não altera saldos"),
        ("D02", "Pública - DFP", "Ausência de nota específica de estoques (R$ 3.109.940 mil; CPC 16 item 36)", "BP 1.01.04", "Julgamento (divulgação)", 0, 0,
         "Não corrigida (comunicar)", "Omissão de divulgação relevante por natureza; saldo > MG"),
        ("D03", "Pública - DFP", "Remissões incorretas: NE 2.16 cita NE 15 para empréstimo em reais (é NE 13) e 'Nota 2929'", "NE 2.16", "Factual (divulgação)",
         0, 0, "Não corrigida", "Qualitativamente irrelevante; deficiência de revisão do relatório"),
        ("D04", "Pública - DFP", "Reclassificação de R$ 29.885 mil entre ações em tesouraria e lucro na venda de ações nos comparativos sem nota", "PL 2.03.02",
         "Factual (apresentação)", 0, 0, "Não corrigida", "Sem efeito no PL total; CPC 23/CPC 26 item 41 requer divulgação"),
        ("D05", "Pública - DFP", "MTM de swap de empréstimo (R$ 11.020 mil) apresentado na linha de debêntures", "BP 2.02.01.02", "Factual (classificação)",
         0, 0, "Não corrigida", "Classificação dentro do mesmo grupo; abaixo da ME"),
        ("D06", "Pública - DFP", "D&A: DFC x DVA (R$ 1.035 mil) e CPV x NE 15 (R$ 5.147 mil)", "DFC/DVA/NE 15", "Factual", 0, 0, "Trivial - não acumular",
         "Abaixo do limite trivial"),
        ("S01", "SIMULADA - base", "E09 Capitalização indevida de OPEX (lançamentos manuais)", "Imobilizado em andamento", "Factual (fraude simulada)",
         -imp.get("E09-Lançamento manual suspeito", 0), -imp.get("E09-Lançamento manual suspeito", 0), "Corrigida (simulação)",
         "Fraude: relevante por natureza independentemente do valor - comunicar à governança (NBC TA 240 item 40)"),
        ("S02", "SIMULADA - base", "E06 Custos digitados x10", "Imobilizado em andamento", "Factual", 0, -imp.get("E06-Digitação x10", 0),
         "Corrigida (simulação)", "Erro"),
        ("S03", "SIMULADA - base", "E01 Duplicidades", "Imobilizado", "Factual", 0, -imp.get("E01-Duplicidade", 0), "Corrigida (simulação)", "Erro"),
        ("S04", "SIMULADA - base", "E02 Corte - NFs de 2026", "Imobilizado em andamento", "Factual", 0, -imp.get("E02-Corte", 0), "Corrigida (simulação)", "Erro de corte"),
        ("S05", "SIMULADA - base", "E03 Depreciação acumulada > custo", "Imobilizado em operação", "Factual", -imp.get("E03-VC negativo", 0),
         -imp.get("E03-VC negativo", 0), "Não corrigida", "Subavaliação"),
        ("S06", "SIMULADA - base", "E04 Depreciação em ativo em andamento", "Imobilizado em andamento", "Factual", -imp.get("E04-Depreciação em andamento", 0),
         -imp.get("E04-Depreciação em andamento", 0), "Não corrigida", "Subavaliação"),
        ("S07", "SIMULADA - base", "E05 Depreciação linear com metade da taxa", "Imobilizado em operação", "Factual", -imp.get("E05-Taxa incorreta", 0),
         -imp.get("E05-Taxa incorreta", 0), "Não corrigida", "Superavaliação"),
    ]
    p0 = l
    for rr in linhas:
        for j, v in enumerate(rr, 1):
            c = ws.cell(l, j, v)
            if j in (6, 7):
                c.number_format = FMT2
            if rr[1].startswith("SIMULADA"):
                c.fill = SINT
        l += 1
    ws.cell(l + 1, 3, "Total das distorções NÃO corrigidas (efeito ativo/PL)").font = B
    ws.cell(l + 1, 7, f'=SUMPRODUCT((H{p0}:H{l-1}="Não corrigida")*G{p0}:G{l-1})').number_format = FMT2
    ws.cell(l + 2, 3, "% da materialidade global").font = B
    ws.cell(l + 2, 7, f"=ABS(G{l+1})/B5").number_format = PCT
    ws.cell(l + 3, 3, "Total corrigido (simulação)").font = B
    ws.cell(l + 3, 7, f'=SUMPRODUCT((H{p0}:H{l-1}="Corrigida (simulação)")*G{p0}:G{l-1})').number_format = FMT2
    ws.cell(l + 4, 3, "Conclusão quantitativa").font = B
    ws.cell(l + 4, 7, f'=IF(ABS(G{l+1})<F5,"Abaixo do trivial",IF(ABS(G{l+1})<D5,"Abaixo da ME","Avaliar efeito na opinião"))')
    ws.cell(l + 6, 3, f"LSD do MUS após ajustes simulados: R$ {r['mus_pos_ajustes']['limite_superior_distorcao_brl']/1e6:,.1f} mi (< distorção tolerável de R$ 93 mi).")

    ws2 = wb.create_sheet("Inseridos_x_Detectados")
    l2 = cabecalho(ws2, "Registro comparativo: erros inseridos x exceções identificadas", "Exigência do item 3 das instruções", sintetico=True)
    df_para_ws(ws2, pd.read_csv(RESULTADOS / "PT450_inseridos_x_detectados.csv"), l2, None, [30, 30, 10, 10, 12, 12, 14, 20, 14])

    ws3 = wb.create_sheet("PT600_Componentes")
    l3 = cabecalho(ws3, "PT 600 - Componentes do grupo e materialidade de componente (R$ mil)", "NBC TA 600 (R2) itens 22-35, A80-A89")
    co = pd.read_csv(RESULTADOS / "PT600_componentes.csv")
    l3 = tabela(ws3, l3, ["Componente", "Receita líquida", "Imobilizado operação", "Intangível", "ARO", "% receita", "% ativo total (imob.+intang.)",
                          "Significativo?", "Materialidade do componente"], [42, 16, 18, 14, 14, 10, 14, 34, 18])
    for rr in co.itertuples(index=False):
        ws3.cell(l3, 1, rr.componente)
        for j, v in enumerate([rr.receita_liquida, rr.imobilizado_operacao_ne8, rr.intangivel_ne9, rr.aro_ne17], 2):
            ws3.cell(l3, j, v).number_format = FMT
        ws3.cell(l3, 6, f"=B{l3}/{conta('DRE', '3.01')}").number_format = PCT
        ws3.cell(l3, 7, f"=(C{l3}+D{l3})/{conta('BPA', '1')}").number_format = PCT
        ws3.cell(l3, 8, f'=IF(OR(F{l3}>0.15,G{l3}>0.15),"Sim - auditoria completa","Não - procedimentos específicos/analíticos")')
        c = ws3.cell(l3, 9, rr.materialidade_componente)
        c.number_format, c.fill = FMT, ENT
        l3 += 1
    ws3.cell(l3 + 1, 1, "Regra: materialidade de componente < materialidade de execução do grupo (93.000); limite agregado coberto pela ME.")
    wb.save(OUT / "PT450_PT600_Distorcoes_Componentes.xlsx")


def cobertura_fontes():
    wb = Workbook()
    ws = wb.active
    ws.title = "Matriz_Cobertura"
    l = cabecalho(ws, "Matriz de cobertura das normas x papéis de trabalho", "NBC TA Estrutura Conceitual e NBC TA 200 a 600")
    df_para_ws(ws, pd.DataFrame(MATRIZ), l, None, [26, 14, 58, 26, 40, 16])
    ws2 = wb.create_sheet("Fontes")
    l2 = cabecalho(ws2, "Relação das fontes utilizadas", "NBC TA 230 / NBC TA 500")
    df_para_ws(ws2, pd.DataFrame(FONTES), l2, None, [6, 58, 22, 70, 14, 18, 30])
    wb.save(OUT / "Matriz_Cobertura_e_Fontes.xlsx")


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    for f in [pt320, pt530, pt520, pt500, pt540_570, pt450_600, cobertura_fontes]:
        f()
        print("ok", f.__name__)
